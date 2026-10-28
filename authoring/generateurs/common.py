"""Cadre commun des générateurs de données — édition « Solen BI 22 h » (V7).

Un script de mission (missionNN_*.py) fait quatre choses :
  1. il décrit la mission dans un MissionSpec (textes, tables, indicateurs, SQL d'analyse) ;
  2. il fabrique la vérité terrain : des événements propres, avec leurs clés métier ;
  3. il rend les sources « sales » (CSV avec anomalies comptées, Excel, JSON, SQLite) ;
  4. il appelle produire_mission() : les valeurs attendues sont calculées DEPUIS LA VÉRITÉ
     (un entrepôt de vérité construit directement, sans lire les sources), jamais depuis un ETL.

Tout est déterministe : une graine par mission, aucune date « maintenant », zip Excel figé.
"""
from __future__ import annotations

import csv
import hashlib
import json
import random
import re
import shutil
import sqlite3
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Callable

# --------------------------------------------------------------------------- chemins
ICI = Path(__file__).resolve().parent
DEPOT = ICI.parents[1]                                    # github/
MODELE = DEPOT / "modele-binome"
FORMATEUR = DEPOT.parent / "formateur" / "fable-v7-solen-22h"
VERITE = FORMATEUR / "solutions" / "verite_terrain"       # privé : vérité terrain brute
V7_DATA = ICI / "v7-data.json"

PALETTE = {"marine": "#2B3180", "deep": "#07152F", "orange": "#FF9A3C",
           "pale": "#F3F5FC", "encre": "#1A1F4D"}
SERIES = ["#2B3180", "#FF9A3C", "#07152F", "#8A90D6", "#C96A12", "#5A5F8F"]
MOIS_FR = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet",
           "août", "septembre", "octobre", "novembre", "décembre"]
DATE_FIXE = datetime(2026, 10, 28, 9, 0, 0)
TAILLE_MAX_MISSION = 1_500_000                  # contrat de l'édition : dossier de mission < 1,5 Mo


def rng_mission(graine: int) -> random.Random:
    """Générateur pseudo-aléatoire à graine fixe (une graine par mission)."""
    return random.Random(graine)


# --------------------------------------------------------------------------- dates
def jours(debut: date, fin: date) -> list[date]:
    """Tous les jours de debut à fin, bornes incluses."""
    return [debut + timedelta(days=i) for i in range((fin - debut).days + 1)]


def date_key(d: date) -> int:
    return d.year * 10000 + d.month * 100 + d.day


def format_date(d: date, style: str = "iso") -> str:
    if style == "iso":
        return d.isoformat()
    if style == "fr":
        return d.strftime("%d/%m/%Y")
    if style == "long":
        return f"{d.day} {MOIS_FR[d.month - 1]} {d.year}"
    raise ValueError(style)


STYLES_DATE = ("fr", "long")


def calendrier(debut: date, fin: date) -> list[dict]:
    """Lignes de dim_date : une par jour calendaire."""
    return [{"date_key": date_key(d), "date_complete": d.isoformat(),
             "annee": d.year, "mois": d.month, "jour": d.day} for d in jours(debut, fin)]


# --------------------------------------------------------------------------- montants
def format_montant(m: int, style: str = "point3") -> str:
    """Rend un montant en millimes sous forme de TND texte.

    point3 « 12.500 » (forme canonique) · virgule3 « 12,500 » · court « 12.5 »
    virgule_court « 12,5 » · milliers « 1 240,000 » (espace = séparateur de milliers).
    """
    if m < 0:
        return "-" + format_montant(-m, style)
    entier, reste = divmod(m, 1000)
    if style == "point3":
        return f"{entier}.{reste:03d}"
    if style == "virgule3":
        return f"{entier},{reste:03d}"
    if style in ("court", "virgule_court"):
        t = f"{entier}.{reste:03d}".rstrip("0").rstrip(".")
        return t.replace(".", ",") if style == "virgule_court" else t
    if style == "milliers":
        groupes = f"{entier:,}".replace(",", " ")
        return f"{groupes},{reste:03d}"
    raise ValueError(style)


STYLES_MONTANT = ("virgule3", "court", "virgule_court", "milliers")


def lire_montant_millimes(texte: str) -> int:
    """Lecture de référence d'un montant TND texte → millimes (sert aux auto-contrôles)."""
    t = str(texte).strip().replace(" ", "").replace(" ", "").replace(",", ".")
    signe = -1 if t.startswith("-") else 1
    t = t.lstrip("+-")
    if not re.fullmatch(r"\d+(\.\d{1,3})?", t):
        raise ValueError(f"montant illisible : {texte!r}")
    entier, _, dec = t.partition(".")
    return signe * (int(entier) * 1000 + int((dec + "000")[:3]))


def millimes(tnd: float) -> int:
    return int(round(tnd * 1000))


def tnd_fr(m: int) -> str:
    """Affichage français : 12 345,678 TND."""
    signe = "-" if m < 0 else ""
    entier, reste = divmod(abs(int(m)), 1000)
    return f"{signe}{entier:,}".replace(",", " ") + f",{reste:03d} TND"


def pct_fr(x: float, dec: int = 2) -> str:
    return f"{x * 100:.{dec}f}".replace(".", ",") + " %"


def _decimal_fr(x: float, dec: int) -> str:
    return f"{x:,.{dec}f}".replace(",", " ").replace(".", ",")


def jours_fr(x: float, dec: int = 2) -> str:
    return _decimal_fr(x, dec) + " jours"


def heures_fr(x: float, dec: int = 2) -> str:
    return _decimal_fr(x, dec) + " h"


# Unités d'indicateur : millimes (TND) · ratio (affiché en %) · nombre · jours / heures (nombres décimaux,
# comparés avec la tolérance du contrat) · centiemes_heure (entier, 750 = 7,5 h).
UNITES_FLOTTANTES = ("ratio", "jours", "heures")
UNITES_DOC = {"millimes": "TND", "jours": "jours, nombre décimal", "heures": "heures, nombre décimal",
              "centiemes_heure": "centièmes d'heure, 750 = 7,5 h"}


def unite_doc(unite: str) -> str:
    """Libellé d'unité dans les documents (ratio et nombre restent écrits tels quels)."""
    return UNITES_DOC.get(unite, unite)


def variante_code(code: str, rng: random.Random) -> str:
    return rng.choice([f" {code.lower()} ", code.lower(), code.capitalize(), f"{code} ", f" {code}"])


# --------------------------------------------------------------------------- anomalies CSV
MOTIFS_STANDARD = {
    "id_conflit": "même identifiant sur deux lignes différentes : les deux lignes sont rejetées",
    "id_manquant": "identifiant vide",
    "cle_inconnue": "code absent du référentiel ou ligne d'origine inconnue",
    "quantite_invalide": "quantité nulle, négative ou non entière",
}
QUANTITES_INVALIDES = ["-2", "0", "1.5", "-1", "2,5", "0", "-3", "0.5"]   # défaut de Doses.quantites_invalides


@dataclass
class ChampsCSV:
    """Décrit les colonnes du CSV de faits pour l'injection d'anomalies."""
    entete: list[str]
    identifiant: str
    quantite: str
    montants: list[str] = field(default_factory=list)
    dates: list[str] = field(default_factory=list)
    codes: list[str] = field(default_factory=list)
    cles_inconnues: dict[str, list[str]] = field(default_factory=dict)


@dataclass
class Doses:
    """Quantités d'anomalies injectées (chaque type est compté dans attendu_*.json)."""
    doublons_exacts: int = 0
    conflits: int = 0          # nombre de PAIRES (2 lignes rejetées par paire)
    manquants: int = 0
    inconnues: int = 0
    quantites: int = 0
    part_montants: float = 0.30
    part_dates: float = 0.25
    part_codes: float = 0.10
    # valeurs écrites dans la colonne quantité des lignes `quantite_invalide` (dans l'ordre de création) ;
    # à remplacer quand « 1.5 » est une valeur valable (poids, heures…)
    quantites_invalides: list[str] = field(default_factory=lambda: list(QUANTITES_INVALIDES))


@dataclass
class LigneSource:
    valeurs: dict[str, str]
    motif: str | None = None      # None = acceptée ; doublon_exact ; deja_charge ; motif de rejet
    origine: str | None = None    # identifiant de l'événement vrai, si la ligne en vient


@dataclass
class BilanCSV:
    lignes: list[LigneSource]
    compteurs: dict[str, int]

    def rejets(self) -> dict[str, int]:
        out: dict[str, int] = {}
        for l in self.lignes:
            if l.motif in MOTIFS_STANDARD:
                out[l.motif] = out.get(l.motif, 0) + 1
        return dict(sorted(out.items()))

    def detail_rejets(self) -> list[dict]:
        return [{"numero_ligne": i + 2, "motif": l.motif, "identifiant": l.valeurs.get("__id", "")}
                for i, l in enumerate(self.lignes) if l.motif in MOTIFS_STANDARD]


def rendre_evenement(ev: dict, champs: ChampsCSV, doses: Doses, rng: random.Random,
                     compteurs: dict | None = None) -> dict[str, str]:
    """Transforme un événement propre en ligne texte, avec formats volontairement mélangés."""
    fmt_m = rng.random() < doses.part_montants
    fmt_d = rng.random() < doses.part_dates
    code_sale = rng.choice(champs.codes) if champs.codes and rng.random() < doses.part_codes else None
    out: dict[str, str] = {}
    for col in champs.entete:
        v = ev[col]
        if col in champs.montants:
            out[col] = format_montant(v, rng.choice(STYLES_MONTANT) if fmt_m else "point3")
        elif col in champs.dates:
            out[col] = format_date(v, rng.choice(STYLES_DATE) if fmt_d else "iso")
        elif col == code_sale:
            out[col] = variante_code(str(v), rng)
        else:
            out[col] = "" if v is None else str(v)
    if compteurs is not None:
        compteurs["formats_montants"] += int(fmt_m and bool(champs.montants))
        compteurs["formats_dates"] += int(fmt_d and bool(champs.dates))
        compteurs["casse_espaces"] += int(code_sale is not None)
    return out


def construire_csv(evenements: list[dict], champs: ChampsCSV, doses: Doses, rng: random.Random,
                   nouvel_id: Callable[[int], str],
                   rejouees: list[dict[str, str]] | None = None) -> BilanCSV:
    """Rend les événements vrais + injecte les anomalies standard, chacune comptée.

    Les lignes piégées sont des lignes EN PLUS de la vérité : la vérité = les lignes acceptées.
    `rejouees` = lignes texte d'un lot précédent, recopiées telles quelles (preuve du rejeu).
    """
    cpt = {"formats_montants": 0, "formats_dates": 0, "casse_espaces": 0}
    idc = champs.identifiant
    lignes = [LigneSource(rendre_evenement(e, champs, doses, rng, cpt), None, e[idc]) for e in evenements]
    pieges: list[LigneSource] = []
    compteur_id = [0]

    def faux_id() -> str:
        compteur_id[0] += 1
        return nouvel_id(compteur_id[0])

    def copie() -> dict:
        return dict(rng.choice(evenements))

    for _ in range(doses.manquants):
        e = copie()
        v = rendre_evenement(e, champs, doses, rng)
        v[idc] = rng.choice(["", "", " "])
        pieges.append(LigneSource(v, "id_manquant"))
    for _ in range(doses.conflits):
        e = copie()
        e[idc] = faux_id()
        e2 = dict(e)
        e2[champs.quantite] = e[champs.quantite] + 1
        for x in (e, e2):
            pieges.append(LigneSource(rendre_evenement(x, champs, doses, rng), "id_conflit"))
    champs_cles = sorted(champs.cles_inconnues)
    for i in range(doses.inconnues):
        e = copie()
        e[idc] = faux_id()
        ch = champs_cles[i % len(champs_cles)]
        e[ch] = rng.choice(champs.cles_inconnues[ch])
        pieges.append(LigneSource(rendre_evenement(e, champs, doses, rng), "cle_inconnue"))
    for i in range(doses.quantites):
        e = copie()
        e[idc] = faux_id()
        v = rendre_evenement(e, champs, doses, rng)
        v[champs.quantite] = doses.quantites_invalides[i % len(doses.quantites_invalides)]
        pieges.append(LigneSource(v, "quantite_invalide"))
    # les deux lignes d'un conflit restent séparées (comme dans un vrai export)
    for p in pieges:
        lignes.insert(rng.randrange(len(lignes) + 1), p)
    deja: set[int] = set()
    for _ in range(doses.doublons_exacts):
        while True:
            i = rng.randrange(len(lignes))
            if lignes[i].motif is None and id(lignes[i]) not in deja:
                break
        deja.add(id(lignes[i]))
        j = rng.randint(i + 1, min(len(lignes), i + 40))
        lignes.insert(j, LigneSource(dict(lignes[i].valeurs), "doublon_exact", lignes[i].origine))
    if rejouees:
        lignes = [LigneSource(dict(v), "deja_charge", v[idc]) for v in rejouees] + lignes
    for l in lignes:                       # pour le détail privé des rejets
        l.valeurs.setdefault("__id", l.valeurs.get(idc, ""))
    bilan = BilanCSV(lignes, {})
    bilan.compteurs = {
        "lignes": len(lignes),
        "lignes_acceptees_nouvelles": sum(1 for l in lignes if l.motif is None),
        "doublons_exacts": sum(1 for l in lignes if l.motif == "doublon_exact"),
        "deja_charges": sum(1 for l in lignes if l.motif == "deja_charge"),
        **cpt,
    }
    return bilan


# --------------------------------------------------------------------------- écrivains
def ecrire_csv(chemin: Path, entete: list[str], lignes: list[dict], sep: str = ";",
               encodage: str = "utf-8") -> None:
    chemin.parent.mkdir(parents=True, exist_ok=True)
    with open(chemin, "w", encoding=encodage, newline="") as f:
        w = csv.writer(f, delimiter=sep, lineterminator="\n")
        w.writerow(entete)
        for l in lignes:
            w.writerow(["" if l.get(c) is None else l.get(c) for c in entete])


@dataclass
class Feuille:
    nom: str
    entete: list[str]
    lignes: list[list]
    titres: list[str] = field(default_factory=list)   # lignes de titre AVANT l'en-tête (+ 1 ligne vide)


def ecrire_excel(chemin: Path, feuilles: list[Feuille]) -> None:
    """Classeur .xlsx (openpyxl), une ou plusieurs feuilles, octets reproductibles."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    chemin.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    for f in feuilles:
        ws = wb.create_sheet(f.nom)
        for t in f.titres:
            ws.append([t])
            ws.cell(ws.max_row, 1).font = Font(bold=True, size=13, color="2B3180")
        if f.titres:
            ws.append([])
        ws.append(f.entete)
        r = ws.max_row
        for c in range(1, len(f.entete) + 1):
            ws.cell(r, c).font = Font(bold=True, color="FFFFFF")
            ws.cell(r, c).fill = PatternFill("solid", fgColor="2B3180")
        for l in f.lignes:
            ws.append(l)
        for i, h in enumerate(f.entete):
            lettre = ws.cell(r, i + 1).column_letter
            ws.column_dimensions[lettre].width = max(12, min(45, len(str(h)) + 4,
                                                             *[len(str(l[i])) + 3 for l in f.lignes]))
    wb.properties.creator = "Solen Digital — données fictives"
    wb.properties.created = DATE_FIXE
    wb.properties.modified = DATE_FIXE
    wb.save(chemin)
    _zip_deterministe(chemin)


def _zip_deterministe(chemin: Path) -> None:
    with zipfile.ZipFile(chemin) as z:
        entrees = [(i.filename, z.read(i.filename)) for i in z.infolist()]
    fixe = DATE_FIXE.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    with zipfile.ZipFile(chemin, "w", zipfile.ZIP_DEFLATED) as z:
        for nom, data in entrees:
            if nom == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*", rb"\g<1>" + fixe, data)
            info = zipfile.ZipInfo(nom, date_time=(2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            z.writestr(info, data)


def ecrire_json(chemin: Path, obj, indent: int = 2) -> None:
    chemin.parent.mkdir(parents=True, exist_ok=True)
    chemin.write_text(json.dumps(obj, ensure_ascii=False, indent=indent) + "\n", encoding="utf-8")


def ecrire_sqlite(chemin: Path, ddl: str, donnees: dict[str, list[dict]]) -> None:
    """Base opérationnelle : schéma réaliste (pas une étoile) + lignes."""
    chemin.parent.mkdir(parents=True, exist_ok=True)
    if chemin.exists():
        chemin.unlink()
    con = sqlite3.connect(chemin)
    con.executescript(ddl)
    for table, lignes in donnees.items():
        if lignes:
            cols = list(lignes[0])
            con.executemany(f"INSERT INTO {table} ({', '.join(cols)}) VALUES ({', '.join('?' * len(cols))})",
                            [tuple(l[c] for c in cols) for l in lignes])
    con.commit()
    con.execute("VACUUM")
    con.close()


# --------------------------------------------------------------------------- spécification
@dataclass
class Colonne:
    nom: str
    type: str                  # INTEGER | TEXT
    description: str
    cle: bool = False          # clé primaire
    ref: str | None = None     # "dim_date.date_key"
    nullable: bool = False


@dataclass
class Table:
    nom: str
    role: str                  # dimension | dimension_scd2 | fait
    description: str
    colonnes: list[Colonne]
    cle_metier: str | None = None

    @property
    def cle(self) -> str:
        return next(c.nom for c in self.colonnes if c.cle)


def table_dim_date() -> Table:
    return Table("dim_date", "dimension", "Calendrier : une ligne par jour de la période chargée.", [
        Colonne("date_key", "INTEGER", "Clé du jour au format AAAAMMJJ (ex. 20260304).", cle=True),
        Colonne("date_complete", "TEXT", "Date ISO AAAA-MM-JJ."),
        Colonne("annee", "INTEGER", "Année."),
        Colonne("mois", "INTEGER", "Mois de 1 à 12."),
        Colonne("jour", "INTEGER", "Jour du mois."),
    ])


@dataclass
class ColonneSource:
    nom: str
    type: str
    description: str
    exemple: str = ""


@dataclass
class Source:
    fichier: str
    type: str                  # csv | excel | json | sqlite
    description: str
    colonnes: list[ColonneSource]
    pieges: list[str] = field(default_factory=list)


@dataclass
class Indicateur:
    code: str
    libelle: str
    unite: str                 # millimes | ratio | nombre | jours | heures | centiemes_heure
    formule: str               # formule exacte en français
    sql: str                   # agrégat SQL sur la vue « mouvements »
    tolerance: float = 0.0
    principal: bool = True     # affiché comme indicateur de la mission


@dataclass
class SCD2:
    table: str
    cle: str
    cle_metier: str
    attribut: str
    faits: list[tuple[str, str, str]]          # (table de fait, colonne clé SCD2, colonne date_key)
    frontieres: list[tuple[str, str]]          # (clé métier, date ISO) à contrôler
    changement: str = ""                       # phrase : ce qui change, quand
    libelle_attribut: str = ""                 # en-tête de l'attribut dans reference/README.md (défaut : attribut)
    precisions: list[str] = field(default_factory=list)   # puces ajoutées au § 7 du cahier


@dataclass
class Lot:
    nom: str                   # lot1 | lot2
    etape: str                 # initial | increment
    debut: date
    fin: date
    description: str


@dataclass
class MissionSpec:
    code: str                  # "tronc-commun" ou "M01"…
    numero: str                # "00", "01"…
    titre: str
    entreprise: str
    secteur: str
    contexte: str
    demande: str               # mail de la direction (texte)
    grain: str
    sources: list[Source]
    tables: list[Table]
    indicateurs: list[Indicateur]
    vue_analyse: str           # SQL : une ligne par mouvement, colonnes d'axes + colonnes additives
    axes: dict[str, str]       # nom d'agrégat → colonne de la vue (ex. {"par_zone": "zone"})
    scd2: SCD2
    lots: list[Lot]
    regles: list[tuple[str, str, str]]           # (anomalie, exemple, règle)
    motifs: dict[str, str]
    ordre_controles: list[str]
    regles_metier: list[str]
    regle_a_clarifier: dict
    variantes: list[str]
    ecrans_app: list[tuple[str, str]]
    textes_app: list[str]
    export_sommes: dict[str, list[str]]
    visuels: list[dict]
    debut_calendrier: date
    graine: int
    sources_rejets: dict[str, str] = field(default_factory=dict)   # nom source → fichier
    # --- compléments des documents (évitent les ajouts locaux dans les scripts de mission) -------------------
    # regle_a_clarifier = {"question", "options", "decision"} (+ "consequence" facultatif : effet dans l'entrepôt)
    # variantes = 3 évolutions ; ecrans_app = 3 écrans (titre, contenu) ; parcours_app = démonstration conseillée.
    # Un texte peut être une fonction (attendus → texte) quand il cite des valeurs de référence.
    parcours_app: str | Callable[[dict], str] = ""
    question_decision: str | Callable[[dict], str] = ""        # la décision attendue de la direction
    regle_mois: str = "(faits et ajustements réunis). Un mouvement est rangé dans le mois **de sa propre date**."
    sections_cahier: list[tuple[str, str | Callable[[dict], str]]] = field(default_factory=list)  # § 13 et suivants
    complement_reference: Callable[[dict], str] | None = None  # texte ajouté à reference/README.md

    @property
    def dossier(self) -> Path:
        return MODELE / "tronc-commun" if self.code == "tronc-commun" else MODELE / "catalogue" / self.code

    def table(self, nom: str) -> Table:
        return next(t for t in self.tables if t.nom == nom)


@dataclass
class LotResultat:
    """Ce que le script de mission sait d'un lot, directement depuis la vérité."""
    lot: Lot
    sources: dict                          # comptes par fichier source
    nettoyage: dict                        # compteurs informatifs (formats normalisés, doublons…)
    rejets: dict[str, dict[str, int]]      # fichier source → motif → nombre
    tables_verite: dict[str, list[dict]]   # état CUMULÉ de l'entrepôt après ce lot (sans dim_date)
    increment: dict | None = None


# --------------------------------------------------------------------------- entrepôt de vérité
def ddl_entrepot(spec: MissionSpec) -> str:
    blocs = []
    for t in spec.tables:
        lignes = []
        for c in t.colonnes:
            s = f"  {c.nom} {c.type}"
            if c.cle:
                s += " PRIMARY KEY"
            elif not c.nullable:
                s += " NOT NULL"
            lignes.append(s)
        for c in t.colonnes:
            if c.ref:
                tt, cc = c.ref.split(".")
                lignes.append(f"  FOREIGN KEY ({c.nom}) REFERENCES {tt}({cc})")
        blocs.append(f"CREATE TABLE {t.nom} (\n" + ",\n".join(lignes) + "\n);")
    return "\n".join(blocs) + "\n"


def entrepot_verite(spec: MissionSpec, tables: dict[str, list[dict]], fin: date) -> sqlite3.Connection:
    """Construit l'étoile directement depuis la vérité terrain (aucune lecture de source)."""
    con = sqlite3.connect(":memory:")
    con.execute("PRAGMA foreign_keys = ON")
    con.executescript(ddl_entrepot(spec))
    donnees = dict(tables)
    donnees.setdefault("dim_date", calendrier(spec.debut_calendrier, fin))
    for t in spec.tables:
        lignes = donnees.get(t.nom, [])
        cols = [c.nom for c in t.colonnes]
        con.executemany(f"INSERT INTO {t.nom} ({', '.join(cols)}) VALUES ({', '.join('?' * len(cols))})",
                        [tuple(l[c] for c in cols) for l in lignes])
    assert not con.execute("PRAGMA foreign_key_check").fetchall(), "vérité incohérente (clé étrangère)"
    return con


def requete_indicateurs(spec_vue: str, indicateurs: list[dict], axe: str | None = None) -> str:
    """Même gabarit que controles/verifier.py (le vérificateur le reconstruit à l'identique)."""
    sel = ", ".join(f"{i['sql']} AS {i['code']}" for i in indicateurs)
    if axe:
        return (f"WITH mouvements AS ({spec_vue}) SELECT {axe} AS cle, {sel} "
                f"FROM mouvements GROUP BY {axe} ORDER BY {axe}")
    return f"WITH mouvements AS ({spec_vue}) SELECT {sel} FROM mouvements"


def _valeur(ind: dict, v):
    if v is None:
        return None
    return round(float(v), 6) if ind["unite"] in UNITES_FLOTTANTES else int(v)


def calculer_attendu(spec: MissionSpec, con: sqlite3.Connection) -> dict:
    inds = [indicateur_dict(i) for i in spec.indicateurs]
    ligne = con.execute(requete_indicateurs(spec.vue_analyse, inds)).fetchone()
    globaux = {i["code"]: _valeur(i, v) for i, v in zip(inds, ligne)}
    agregats = {}
    for nom, axe in spec.axes.items():
        agregats[nom] = {}
        for r in con.execute(requete_indicateurs(spec.vue_analyse, inds, axe)):
            agregats[nom][str(r[0])] = {i["code"]: _valeur(i, v) for i, v in zip(inds, r[1:])}
    s = spec.scd2
    versions = [dict(zip(("cle_metier", "attribut", "valid_from", "valid_to"), r)) for r in con.execute(
        f"SELECT {s.cle_metier}, {s.attribut}, valid_from, valid_to FROM {s.table} "
        f"ORDER BY {s.cle_metier}, valid_from")]
    frontieres = []
    for cm, d in s.frontieres:
        r = con.execute(f"SELECT {s.attribut} FROM {s.table} WHERE {s.cle_metier} = ? AND valid_from <= ? "
                        f"AND (valid_to IS NULL OR valid_to > ?)", (cm, d, d)).fetchall()
        assert len(r) == 1
        frontieres.append({"cle_metier": cm, "date": d, "attribut": r[0][0]})
    return {
        "tables": {t.nom: con.execute(f"SELECT COUNT(*) FROM {t.nom}").fetchone()[0] for t in spec.tables},
        "indicateurs": globaux,
        "agregats": agregats,
        "scd2": {"table": s.table, "cle_metier": s.cle_metier, "attribut": s.attribut,
                 "versions": versions, "frontieres": frontieres},
    }


def indicateur_dict(i: Indicateur) -> dict:
    return {"code": i.code, "libelle": i.libelle, "unite": i.unite, "formule": i.formule,
            "sql": i.sql, "tolerance": i.tolerance, "principal": i.principal}


def cle_version(versions: list[dict], cle_metier_col: str, cle_col: str, valeur: str, jour: date) -> int:
    """Clé technique de la version SCD2 valable à `jour` (valid_from inclus, valid_to exclu)."""
    d = jour.isoformat()
    ok = [v[cle_col] for v in versions if v[cle_metier_col] == valeur and v["valid_from"] <= d
          and (v["valid_to"] is None or d < v["valid_to"])]
    assert len(ok) == 1, (valeur, d, ok)
    return ok[0]


# --------------------------------------------------------------------------- contrat entrepôt
def contrat_entrepot(spec: MissionSpec) -> dict:
    tables = {}
    for t in spec.tables:
        tables[t.nom] = {
            "role": t.role, "cle": t.cle, "cle_metier": t.cle_metier,
            "colonnes": [c.nom for c in t.colonnes],
            "colonnes_entieres": [c.nom for c in t.colonnes if c.type == "INTEGER"],
            "colonnes_nullables": [c.nom for c in t.colonnes if c.nullable],
            "cles_etrangeres": {c.nom: c.ref for c in t.colonnes if c.ref},
        }
    s = spec.scd2
    return {
        "version_contrat": 1,
        "mission": spec.code, "entreprise": spec.entreprise, "grain": spec.grain,
        "unite_monetaire": "millimes entiers (1 TND = 1 000 millimes)",
        "fichiers": {"entrepot": "entrepot.sqlite", "rejets": "rejets.csv", "export": "export",
                     "manifest": "export/manifest.json", "app": "app/index.html"},
        "ordre_tables": [t.nom for t in spec.tables],
        "tables": tables,
        "faits": [t.nom for t in spec.tables if t.role == "fait"],
        "dimension_date": {"table": "dim_date", "cle": "date_key", "date": "date_complete",
                           "debut": spec.debut_calendrier.isoformat(),
                           "regle": "une ligne par jour, du " + spec.debut_calendrier.isoformat()
                                    + " au dernier jour du dernier lot chargé"},
        "scd2": {"table": s.table, "cle": s.cle, "cle_metier": s.cle_metier, "attribut": s.attribut,
                 "valid_from": "valid_from", "valid_to": "valid_to",
                 "regle": "valid_from inclus, valid_to exclu, valid_to NULL = version courante",
                 "faits": [{"table": a, "colonne": b, "date_key": c} for a, b, c in s.faits],
                 "frontieres": [{"cle_metier": a, "date": b} for a, b in s.frontieres]},
        "vue_analyse": " ".join(spec.vue_analyse.split()),
        "indicateurs": [indicateur_dict(i) for i in spec.indicateurs],
        "axes": spec.axes,
        "rejets": {"colonnes_obligatoires": ["lot", "source", "numero_ligne", "motif"],
                   "motifs": list(spec.motifs), "motifs_toleres": ["doublon_exact"],
                   "sources": sorted(set(spec.sources_rejets.values()))},
        "export": {"tables": [t.nom for t in spec.tables], "sommes": spec.export_sommes,
                   "format": "CSV UTF-8, séparateur virgule, en-tête = noms de colonnes"},
        "app": {"textes_obligatoires": spec.textes_app, "taille_max_octets": 5_000_000},
        "lots": {l.nom: {"etape": l.etape, "attendu": f"attendu_{l.nom}.json",
                         "debut": l.debut.isoformat(), "fin": l.fin.isoformat()} for l in spec.lots},
        "etapes": {"initial": spec.lots[0].nom, "increment": spec.lots[-1].nom, "rejeu": spec.lots[-1].nom},
    }


# --------------------------------------------------------------------------- graphiques
def _plt():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams.update({"svg.hashsalt": "solen-v7", "font.family": "DejaVu Sans", "font.size": 10,
                         "axes.edgecolor": PALETTE["encre"], "axes.labelcolor": PALETTE["encre"],
                         "xtick.color": PALETTE["encre"], "ytick.color": PALETTE["encre"],
                         "axes.spines.top": False, "axes.spines.right": False})
    return plt


def _nombre_fr(x, _pos=None):
    if abs(x - round(x)) < 1e-9:
        return f"{int(round(x)):,}".replace(",", " ")
    return f"{x:,.1f}".replace(",", " ").replace(".", ",")


def _sauver(fig, base: Path) -> None:
    base.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(base.with_suffix(".png"), dpi=150, bbox_inches="tight", metadata={"Software": None})
    fig.savefig(base.with_suffix(".svg"), bbox_inches="tight", metadata={"Date": None, "Creator": None})


def _echelle(ind: dict):
    return {"millimes": (1000.0, "TND"), "ratio": (0.01, "%"), "jours": (1.0, "jours"), "heures": (1.0, "h"),
            "centiemes_heure": (100.0, "h")}.get(ind["unite"], (1.0, ""))


def _titre_axe(libelle: str, unite: str) -> str:
    """« Coût (TND) » ; un comptage n'a pas d'unité : « Interventions », jamais « Interventions () »."""
    return f"{libelle} ({unite})" if unite else libelle


def _legende_dessous(ax, n: int, titre: str | None = None) -> None:
    """Légende sous le graphique : elle ne recouvre jamais une barre."""
    ax.legend(title=titre, frameon=False, ncol=min(max(n, 1), 5), loc="upper center", bbox_to_anchor=(0.5, -0.14))


def visuel_empile(spec: MissionSpec, con, v: dict, base: Path) -> None:
    from matplotlib.ticker import FuncFormatter
    plt = _plt()
    ind = indicateur_dict(next(i for i in spec.indicateurs if i.code == v["indicateur"]))
    q = (f"WITH mouvements AS ({spec.vue_analyse}) SELECT {v['x']}, {v['serie']}, {ind['sql']} "
         f"FROM mouvements GROUP BY 1, 2 ORDER BY 1, 2")
    rows = con.execute(q).fetchall()
    xs = sorted({r[0] for r in rows})
    series = list(dict.fromkeys(v.get("ordre_series") or sorted({r[1] for r in rows})))
    div, unite = _echelle(ind)
    fig, ax = plt.subplots(figsize=(8.5, 4.8))
    bas = [0.0] * len(xs)
    for k, s in enumerate(series):
        val = [(next((r[2] for r in rows if r[0] == x and r[1] == s), 0) or 0) / div for x in xs]
        ax.bar(xs, val, bottom=bas, label=s, color=SERIES[k % len(SERIES)], width=0.65)
        bas = [a + b for a, b in zip(bas, val)]
    ax.set_title(v["titre"], color=PALETTE["deep"], fontweight="bold", loc="left")
    ax.set_ylabel(_titre_axe(v.get("libelle_y", ind["libelle"]), unite))
    ax.set_xlabel(v.get("libelle_x", v["x"]))
    ax.yaxis.set_major_formatter(FuncFormatter(_nombre_fr))
    _legende_dessous(ax, len(series), v.get("libelle_serie", v["serie"]))
    ax.grid(axis="y", color=PALETTE["pale"])
    ax.set_axisbelow(True)
    _sauver(fig, base)
    plt.close(fig)


def visuel_barres(spec: MissionSpec, con, v: dict, base: Path) -> None:
    """Barres par valeur d'un axe, un ou plusieurs indicateurs de même unité.

    Options : `top` (N valeurs les plus hautes du premier indicateur), `horizontal` (barres couchées, triées,
    lisibles avec beaucoup d'étiquettes ; valeurs négatives en orange), `libelle_y`, `libelle_x`.
    Au-delà de 8 étiquettes verticales, elles sont inclinées.
    """
    from matplotlib.ticker import FuncFormatter
    plt = _plt()
    inds = [indicateur_dict(next(i for i in spec.indicateurs if i.code == c)) for c in v["indicateurs"]]
    q = requete_indicateurs(spec.vue_analyse, inds, v["axe"])
    rows = con.execute(q).fetchall()
    if v.get("top"):
        rows = sorted(rows, key=lambda r: (-(r[1] or 0), str(r[0])))[:v["top"]]
    div, unite = _echelle(inds[0])
    dec = {"TND": 0, "": 0, "%": 1}.get(unite, 2)
    n = len(inds)
    titre_val = _titre_axe(v.get("libelle_y", inds[0]["libelle"]), unite)
    if v.get("horizontal"):
        rows = sorted(rows, key=lambda r: ((r[1] or 0), str(r[0])))
        fig, ax = plt.subplots(figsize=(8.5, max(4.6, 0.32 * len(rows) * n + 1.2)))
    else:
        fig, ax = plt.subplots(figsize=(8.5, 4.8))
    larg = 0.8 / n
    haut = 0.0
    for k, ind in enumerate(inds):
        pos = [i + (k - (n - 1) / 2) * larg for i in range(len(rows))]
        vals = [(r[1 + k] or 0) / div for r in rows]
        haut = max([haut] + [abs(x) for x in vals])
        couleur = ([PALETTE["orange"] if x < 0 else SERIES[0] for x in vals] if n == 1 and v.get("horizontal")
                   else SERIES[k % len(SERIES)])
        trace = ax.barh if v.get("horizontal") else ax.bar
        barres = trace(pos, vals, larg, label=ind["libelle"], color=couleur)
        ax.bar_label(barres, labels=[_nombre_fr(round(x)) if dec == 0 else
                                     f"{x:,.{dec}f}".replace(",", " ").replace(".", ",") for x in vals],
                     fontsize=8 if len(rows) * n <= 12 else 7, color=PALETTE["encre"], padding=2)
    etiquettes = [str(r[0]) for r in rows]
    if v.get("horizontal"):
        ax.set_yticks(range(len(rows)), etiquettes)
        ax.set_xlabel(titre_val)
        ax.set_ylabel(v.get("libelle_x", v["axe"]))
        ax.xaxis.set_major_formatter(FuncFormatter(_nombre_fr))
        ax.margins(x=0.15)
        ax.grid(axis="x", color=PALETTE["pale"])
        if any((r[1] or 0) < 0 for r in rows):
            ax.axvline(0, color=PALETTE["encre"], linewidth=0.8)
    else:
        incline = len(rows) > 8
        ax.set_xticks(range(len(rows)), etiquettes, rotation=45 if incline else 0, ha="right" if incline else "center")
        ax.set_ylabel(titre_val)
        ax.set_xlabel(v.get("libelle_x", v["axe"]))
        ax.yaxis.set_major_formatter(FuncFormatter(_nombre_fr))
        ax.set_ylim(top=haut * 1.15 if haut else 1)
        ax.grid(axis="y", color=PALETTE["pale"])
    ax.set_title(v["titre"], color=PALETTE["deep"], fontweight="bold", loc="left")
    if n > 1:
        _legende_dessous(ax, n)
    ax.set_axisbelow(True)
    _sauver(fig, base)
    plt.close(fig)


def visuel_rejets(attendus: dict[str, dict], v: dict, base: Path) -> None:
    plt = _plt()
    motifs = sorted({m for a in attendus.values() for m in a["rejets"]["par_motif"]})
    fig, ax = plt.subplots(figsize=(8.5, 4.2))
    n = len(attendus)
    larg = 0.8 / n
    for k, (lot, a) in enumerate(attendus.items()):
        xs = [i + (k - (n - 1) / 2) * larg for i in range(len(motifs))]
        vals = [a["rejets"]["par_motif"].get(m, 0) for m in motifs]
        b = ax.bar(xs, vals, width=larg, label=lot, color=SERIES[k % len(SERIES)])
        ax.bar_label(b, fontsize=8, padding=2)
    ax.set_xticks(range(len(motifs)), motifs)
    ax.set_title(v["titre"], color=PALETTE["deep"], fontweight="bold", loc="left")
    ax.set_ylabel("Lignes rejetées")
    ax.set_xlabel("Motif de rejet")
    ax.set_ylim(top=max([1] + [x for a in attendus.values() for x in a["rejets"]["par_motif"].values()]) * 1.15)
    _legende_dessous(ax, n)
    ax.grid(axis="y", color=PALETTE["pale"])
    ax.set_axisbelow(True)
    _sauver(fig, base)
    plt.close(fig)


# --------------------------------------------------------------------------- documents
def md_tableau(entete: list[str], lignes: list[list]) -> str:
    out = ["| " + " | ".join(entete) + " |", "|" + "|".join("---" for _ in entete) + "|"]
    out += ["| " + " | ".join(str(x).replace("|", "\\|") for x in l) + " |" for l in lignes]
    return "\n".join(out)


def _ind_valeur(i: dict, v) -> str:
    if v is None:
        return "—"
    u = i["unite"]
    if u == "ratio":
        return pct_fr(v)
    if u == "millimes":
        return tnd_fr(v)
    if u == "jours":
        return jours_fr(v)
    if u == "heures":
        return heures_fr(v)
    if u == "centiemes_heure":
        return heures_fr(v / 100) + f" ({v} centièmes)"
    return f"{v}"


def _texte(t, attendus: dict | None) -> str:
    """Texte de MissionSpec : chaîne, ou fonction des valeurs attendues (attendus par lot)."""
    return t(attendus or {}) if callable(t) else (t or "")


def _liste_numerotee(elements) -> str:
    return "\n".join(f"{k}. {e}" for k, e in enumerate(elements, 1))


def _ecrans(spec: MissionSpec) -> str:
    return _liste_numerotee(f"**{t}** — {c}" for t, c in spec.ecrans_app)


def doc_fiche(spec: MissionSpec, attendus: dict | None = None) -> str:
    inds = "\n".join(f"- **{i.libelle}** ({unite_doc(i.unite)}) : {i.formule}"
                     for i in spec.indicateurs if i.principal)
    ecrans = _ecrans(spec)
    variantes = _liste_numerotee(spec.variantes)
    q = spec.regle_a_clarifier
    decision = _texte(spec.question_decision, attendus)
    decision = f"## La décision attendue\n\n{decision}\n\n" if decision else ""
    lots = "\n".join(f"- `{l.nom}` ({l.debut.isoformat()} → {l.fin.isoformat()}) : {l.description}" for l in spec.lots)
    return f"""# Mission {spec.numero} — {spec.titre}

**Client fictif :** {spec.entreprise} ({spec.secteur}). **Grain du fait :** {spec.grain}.
Données 100 % synthétiques : noms, chiffres et personnes sont inventés.

## Le contexte

{spec.contexte}

## Le message de la direction

> {spec.demande.replace(chr(10), chr(10) + '> ')}

## Ce que vous livrez

1. `outputs/rejets.csv` : chaque ligne écartée avec son motif.
2. `outputs/entrepot.sqlite` : une étoile ({', '.join(t.nom for t in spec.tables)}).
3. Le lot 2 chargé, l'historique gardé et un rejeu sans doublon. {spec.scd2.changement}
4. `outputs/export/*.csv` + `manifest.json`, puis le rapport Power BI actualisé.
5. `outputs/app/index.html` : une mini-application locale de trois écrans.
6. Une recommandation courte à la direction, appuyée sur les chiffres.

{decision}## Les indicateurs (définitions exactes)

{inds}

Tous les montants sont stockés en **millimes entiers** (1 TND = 1 000 millimes) et affichés en TND à 3 décimales.

## Les données

{lots}

Un `echantillon/` de 12 lignes permet de tout calculer à la main d'abord (voir `echantillon/attendu.md`).

## La règle à clarifier avec la direction

**Question :** {q['question']}

{q['options']}

**Décision de la direction pour cette mission :** {q['decision']}

## Les trois écrans de l'application

{ecrans}

## Trois variantes possibles (évolutions)

{variantes}

## Comment savoir que c'est juste ?

```bash
python -m controles.verifier --mission {spec.code} --etape tout
python -m pytest
```

Les valeurs attendues sont dans `reference/attendu_lot1.json` et `reference/attendu_lot2.json`. Elles ont été calculées par le générateur depuis la vérité, pas par une solution.
"""


def doc_cahier(spec: MissionSpec, attendus: dict | None = None) -> str:
    srcs = md_tableau(["Fichier", "Type", "Contenu", "Pièges connus"],
                      [[f"`{s.fichier}`", s.type, s.description, " ; ".join(s.pieges) or "—"] for s in spec.sources])
    regles = md_tableau(["Anomalie", "Exemple", "Règle"], [list(r) for r in spec.regles])
    motifs = md_tableau(["Motif (`motif`)", "Signification"], [[f"`{k}`", v] for k, v in spec.motifs.items()])
    ordre = "\n".join(f"{k}. {o}" for k, o in enumerate(spec.ordre_controles, 1))
    metier = "\n".join(f"- {r}" for r in spec.regles_metier)
    tables = []
    for t in spec.tables:
        cols = md_tableau(["Colonne", "Type", "Règle"], [[f"`{c.nom}`", c.type + (" (clé)" if c.cle else "")
                                                           + (f" → `{c.ref}`" if c.ref else ""), c.description]
                                                          for c in t.colonnes])
        tables.append(f"### `{t.nom}` — {t.role.replace('_', ' ')}\n\n{t.description}\n\n{cols}\n")
    inds = md_tableau(["Indicateur", "Unité", "Formule exacte"],
                      [[i.libelle, "TND (stocké en millimes)" if i.unite == "millimes" else unite_doc(i.unite), i.formule]
                       for i in spec.indicateurs])
    s = spec.scd2
    precisions = "".join(f"\n- {p}" for p in s.precisions)
    sommes = "; ".join(f"`{t}` : {', '.join(c)}" for t, c in spec.export_sommes.items())
    csv_faits = next((x.fichier for x in spec.sources if x.type == "csv"), spec.sources[0].fichier)
    q = spec.regle_a_clarifier
    # § 12 règle à clarifier, puis § 13… : sections propres à la mission, variantes, parcours, décision attendue
    fin = [("Trois variantes (évolutions possibles)", _liste_numerotee(spec.variantes)),
           ("Parcours de démonstration de l'application (3 écrans)",
            _ecrans(spec) + (f"\n\n{_texte(spec.parcours_app, attendus)}" if spec.parcours_app else ""))]
    if spec.question_decision:
        fin.append(("La décision attendue", _texte(spec.question_decision, attendus)))
    suite = "".join(f"\n## {k}. {t}\n\n{_texte(x, attendus)}\n"
                    for k, (t, x) in enumerate(list(spec.sections_cahier) + fin, 13))
    consequence = f"\n\n{q['consequence']}" if q.get("consequence") else ""
    return f"""# Cahier des charges — Mission {spec.numero} : {spec.titre}

Client fictif : **{spec.entreprise}**. Ce document fixe les règles. En cas de doute, il gagne sur toute autre explication. Les noms de tables et de colonnes ci-dessous sont **obligatoires** : les contrôles automatiques et le guide Power BI les utilisent.

## 1. Grain

**Une ligne de fait = {spec.grain}.** Chaque fait a une clé métier unique (identifiant texte de la source).

## 2. Sources (identiques dans `lot1/` et `lot2/`)

Les sources sont **lues, jamais modifiées**.

{srcs}

## 3. Règles de nettoyage

{regles}

## 4. Motifs de rejet

{motifs}

Ordre des contrôles sur chaque ligne (le premier problème trouvé donne le motif) :

{ordre}

`outputs/rejets.csv` : une ligne par ligne source rejetée, en UTF-8, avec au moins les colonnes `lot` (`lot1` ou `lot2`), `source` (nom du fichier, ex. `{csv_faits}`), `numero_ligne` (numéro de ligne dans le fichier, l'en-tête étant la ligne 1 ; pour la base SQLite, l'identifiant de la ligne), `motif`. Le fichier garde les rejets de tous les lots ; relancer un lot **remplace** ses rejets (pas de doublon). Les doublons exacts ne sont pas des rejets : on garde une occurrence (on peut les noter avec le motif `doublon_exact`, facultatif).

## 5. Règles métier

{metier}

## 6. Entrepôt `outputs/entrepot.sqlite` (modèle en étoile)

Charger les dimensions **avant** les faits. Clés techniques `*_key` entières ; clés métier `*_id` en texte. Montants en **millimes entiers** (type INTEGER). Les clés étrangères sont déclarées (`FOREIGN KEY`).

{chr(10).join(tables)}

## 7. Historique (SCD2) sur `{s.table}`

{s.changement}

- Une ligne par version : `valid_from` **inclus**, `valid_to` **exclu**, `valid_to` vide (NULL) = version courante.
- La première version commence à la date `valide_depuis` donnée par la source.
- Une nouvelle version **ferme** la précédente : `valid_to` de l'ancienne = `valid_from` de la nouvelle.
- Un fait est rattaché à la version valable **à la date du fait** : un fait ancien garde l'ancienne version.{precisions}

## 8. Chargement incrémental et rejeu

- `initial` : entrepôt neuf + `lot1`.
- `increment` : ajoute `lot2` sans effacer `lot1` (nouvelles lignes, nouvelle version SCD2, calendrier prolongé).
- `rejeu` : relancer `lot2` ne change **rien** (aucune ligne ajoutée, mêmes rejets). Un fait dont la clé métier est déjà dans l'entrepôt est ignoré.
- `lot2` contient volontairement des lignes de `lot1` (chevauchement d'export) : elles ne créent pas de doublon.

## 9. Indicateurs (définitions exactes)

{inds}

Le calcul de référence est fait sur la vue « mouvements » de `reference/contrat_entrepot.json` {spec.regle_mois}

## 10. Sorties

- `outputs/rejets.csv` (voir § 4).
- `outputs/export/<table>.csv` pour chaque table de l'entrepôt (UTF-8, séparateur virgule, en-tête = noms de colonnes) + `outputs/export/manifest.json` : `{{"tables": {{"<table>": {{"fichier": "<table>.csv", "lignes": N, "sommes": {{"<colonne>": S}}}}}}}}`. Sommes obligatoires : {sommes}.
- `outputs/app/index.html` : **un seul fichier** autonome (CSS, JavaScript et données à l'intérieur, aucun lien Internet), avec les écrans {", ".join(f"« {t} »" for t in spec.textes_app)} et des filtres.

## 11. Contrôles

```bash
python -m controles.verifier --mission {spec.code} --etape initial     # après le lot 1
python -m controles.verifier --mission {spec.code} --etape increment   # après le lot 2
python -m controles.verifier --mission {spec.code} --etape rejeu       # après avoir relancé le lot 2
python -m controles.verifier --mission {spec.code} --etape export
python -m controles.verifier --mission {spec.code} --etape tout
```

## 12. Règle à clarifier

{q['question']} — **Décision :** {q['decision']}

Options étudiées :

{q['options']}{consequence}
{suite}"""


def doc_dictionnaire(spec: MissionSpec) -> str:
    parts = [f"# Dictionnaire des données — Mission {spec.numero} ({spec.entreprise})\n",
             "Ce dictionnaire décrit chaque colonne des sources, puis chaque colonne de l'entrepôt.\n",
             "## Sources\n"]
    for s in spec.sources:
        parts.append(f"### `{s.fichier}` ({s.type})\n\n{s.description}\n")
        parts.append(md_tableau(["Colonne", "Type", "Description", "Exemple"],
                                [[f"`{c.nom}`", c.type, c.description, f"`{c.exemple}`" if c.exemple else ""]
                                 for c in s.colonnes]) + "\n")
    parts.append("## Entrepôt\n")
    for t in spec.tables:
        parts.append(f"### `{t.nom}`\n\n{t.description}\n")
        parts.append(md_tableau(["Colonne", "Type", "Description"],
                                [[f"`{c.nom}`", c.type, c.description] for c in t.colonnes]) + "\n")
    parts.append("## Indicateurs\n")
    parts.append(md_tableau(["Code", "Libellé", "Unité", "Formule"],
                            [[f"`{i.code}`", i.libelle, i.unite if i.unite == "millimes" else unite_doc(i.unite),
                              i.formule] for i in spec.indicateurs]) + "\n")
    return "\n".join(parts)


def doc_reference_readme(spec: MissionSpec, attendus: dict[str, dict], visuels: list[str]) -> str:
    inds = [indicateur_dict(i) for i in spec.indicateurs]
    lignes = [[i["libelle"]] + [_ind_valeur(i, a["indicateurs"][i["code"]]) for a in attendus.values()]
              + [f"`{i['code']}`"] for i in inds]
    t_ind = md_tableau(["Indicateur"] + [f"après {l}" for l in attendus] + ["code"], lignes)
    t_tab = md_tableau(["Table"] + [f"après {l}" for l in attendus],
                       [[f"`{t}`"] + [a["tables"][t] for a in attendus.values()]
                        for t in next(iter(attendus.values()))["tables"]])
    motifs = sorted({m for a in attendus.values() for m in a["rejets"]["par_motif"]})
    t_rej = md_tableau(["Motif"] + list(attendus), [[f"`{m}`"] + [a["rejets"]["par_motif"].get(m, 0)
                                                                  for a in attendus.values()] for m in motifs]
                       + [["**total**"] + [a["rejets"]["total"] for a in attendus.values()]])
    dern = list(attendus.values())[-1]
    t_scd = md_tableau(["Clé métier", spec.scd2.libelle_attribut or spec.scd2.attribut, "valid_from", "valid_to"],
                       [[v["cle_metier"], v["attribut"], v["valid_from"], v["valid_to"] or "(courante)"]
                        for v in dern["scd2"]["versions"]])
    vis = "\n".join(f"- `visuels/{v}.png` · `visuels/{v}.svg`" for v in visuels)
    return f"""# Référence — Mission {spec.numero} ({spec.entreprise})

Valeurs calculées par le **générateur** (`authoring/generateurs/`) depuis la vérité terrain, sans passer par un ETL. C'est l'oracle de `controles/verifier.py`. Montants en millimes dans les JSON ; ici affichés en TND.

## Indicateurs

{t_ind}

## Lignes par table

{t_tab}

## Rejets par motif (lignes)

{t_rej}

## Historique ({spec.scd2.table}) après le dernier lot

{t_scd}

## Fichiers

- `attendu_lot1.json` : état attendu après `initial`.
- `attendu_lot2.json` : état attendu après `increment` **et** après `rejeu` (identique).
- `contrat_entrepot.json` : noms des tables et colonnes, SCD2, vue d'analyse, indicateurs, axes — lu par le vérificateur.
{vis}
""" + (spec.complement_reference(attendus) if spec.complement_reference else "")


# --------------------------------------------------------------------------- orchestration
def taille_dossier(d: Path) -> int:
    return sum(p.stat().st_size for p in d.rglob("*") if p.is_file())


def empreinte_dossier(d: Path) -> str:
    h = hashlib.sha256()
    for p in sorted(x for x in d.rglob("*") if x.is_file()):
        h.update(str(p.relative_to(d)).encode())
        h.update(p.read_bytes())
    return h.hexdigest()


def preparer_dossier(spec: MissionSpec) -> Path:
    """Vide les sous-dossiers générés de la mission (jamais le reste du dépôt)."""
    d = spec.dossier
    for sous in ("echantillon", "sources", "reference"):
        if (d / sous).exists():
            shutil.rmtree(d / sous)
    d.mkdir(parents=True, exist_ok=True)
    return d


def produire_reference(spec: MissionSpec, lots: list[LotResultat]) -> dict[str, dict]:
    """Calcule et écrit reference/ (attendus, contrat, visuels, README) + documents de la mission."""
    ref = spec.dossier / "reference"
    attendus: dict[str, dict] = {}
    cons = {}
    for lr in lots:
        con = entrepot_verite(spec, lr.tables_verite, lr.lot.fin)
        cons[lr.lot.nom] = con
        calc = calculer_attendu(spec, con)
        rejets_src = {s: dict(sorted(m.items())) for s, m in lr.rejets.items() if m}
        par_motif: dict[str, int] = {}
        for m in rejets_src.values():
            for k, v in m.items():
                par_motif[k] = par_motif.get(k, 0) + v
        a = {
            "mission": spec.code, "entreprise": spec.entreprise, "lot": lr.lot.nom, "etape": lr.lot.etape,
            "periode": {"debut": lr.lot.debut.isoformat(), "fin": lr.lot.fin.isoformat()},
            "calcule_par": "générateur, depuis la vérité terrain (aucun ETL)",
            "unite_monetaire": "millimes entiers (1 TND = 1 000 millimes)",
            "sources": lr.sources,
            "nettoyage": lr.nettoyage,
            "rejets": {"total": sum(par_motif.values()), "par_motif": dict(sorted(par_motif.items())),
                       "par_source": rejets_src},
            **calc,
        }
        if lr.increment is not None:
            a["increment"] = lr.increment
        attendus[lr.lot.nom] = a
    dernier = spec.lots[-1].nom
    attendus[dernier]["rejeu"] = {
        "etape": "rejeu", "lignes_ajoutees": 0,
        "regle": f"relancer {dernier} donne exactement ce fichier : mêmes tables, mêmes rejets, mêmes agrégats"}
    for nom, a in attendus.items():
        ecrire_json(ref / f"attendu_{nom}.json", a)
    ecrire_json(ref / "contrat_entrepot.json", contrat_entrepot(spec))
    noms_visuels = []
    for v in spec.visuels:
        base = ref / "visuels" / v["fichier"]
        if v["type"] == "empile":
            visuel_empile(spec, cons[dernier], v, base)
        elif v["type"] == "barres":
            visuel_barres(spec, cons[dernier], v, base)
        elif v["type"] == "rejets":
            visuel_rejets(attendus, v, base)
        elif v["type"] == "fonction":            # graphique propre à la mission : dessiner(spec, con, attendus, base)
            v["dessiner"](spec, cons[dernier], attendus, base)
        noms_visuels.append(v["fichier"])
    (ref / "README.md").write_text(doc_reference_readme(spec, attendus, noms_visuels), encoding="utf-8")
    d = spec.dossier
    (d / "fiche_mission.md").write_text(doc_fiche(spec, attendus), encoding="utf-8")
    (d / "cahier_des_charges.md").write_text(doc_cahier(spec, attendus), encoding="utf-8")
    (d / "dictionnaire.md").write_text(doc_dictionnaire(spec), encoding="utf-8")
    return attendus


def calculer_attendu_echantillon(spec: MissionSpec, tables: dict[str, list[dict]], debut: date,
                                 fin: date, rejets: dict[str, dict[str, int]], extra: dict) -> dict:
    con = sqlite3.connect(":memory:")
    con.execute("PRAGMA foreign_keys = ON")
    con.executescript(ddl_entrepot(spec))
    t = dict(tables)
    t["dim_date"] = calendrier(debut, fin)
    for tb in spec.tables:
        cols = [c.nom for c in tb.colonnes]
        con.executemany(f"INSERT INTO {tb.nom} ({', '.join(cols)}) VALUES ({', '.join('?' * len(cols))})",
                        [tuple(l[c] for c in cols) for l in t.get(tb.nom, [])])
    calc = calculer_attendu(spec, con)
    par_motif: dict[str, int] = {}
    for m in rejets.values():
        for k, v in m.items():
            par_motif[k] = par_motif.get(k, 0) + v
    return {"mission": spec.code, "echantillon": True, **extra,
            "rejets": {"total": sum(par_motif.values()), "par_motif": dict(sorted(par_motif.items())),
                       "par_source": rejets},
            "tables": calc["tables"], "indicateurs": calc["indicateurs"],
            "agregats": calc["agregats"]}


def bloc_v7_data(spec: MissionSpec, echantillon: dict, attendus: dict[str, dict]) -> dict:
    """Bloc d'une mission dans v7-data.json : vérités de l'échantillon et nombres clés de la référence."""
    inds = [indicateur_dict(i) for i in spec.indicateurs]
    return {
        "numero": spec.numero, "titre": spec.titre, "entreprise": spec.entreprise, "grain": spec.grain,
        "indicateurs": [{k: i[k] for k in ("code", "libelle", "unite", "formule")} for i in inds],
        "echantillon": echantillon,
        "reference": {lot: {"indicateurs": a["indicateurs"], "tables": a["tables"],
                            "rejets": a["rejets"]["par_motif"], "rejets_total": a["rejets"]["total"],
                            "nettoyage": a["nettoyage"], "scd2_versions": a["scd2"]["versions"],
                            "scd2_frontieres": a["scd2"]["frontieres"],
                            "par_axe": a["agregats"]}
                      for lot, a in attendus.items()},
    }


def exporter_v7_data(spec: MissionSpec, echantillon: dict, attendus: dict[str, dict]) -> None:
    """Fusionne le bloc de la mission dans v7-data.json, clé = code de mission (lu par le générateur de schémas).

    Seul le bloc de cette mission est remplacé : les autres missions restent intactes (aucun écrasement),
    quel que soit l'ordre de génération. Écriture atomique (fichier temporaire puis renommage).
    """
    data = json.loads(V7_DATA.read_text(encoding="utf-8")) if V7_DATA.exists() else {}
    data["_lisez_moi"] = ("Généré par authoring/generateurs (un bloc par mission, fusionné par code). Montants en "
                          "millimes. Ne pas éditer à la main : relancer le générateur de la mission.")
    missions = dict(data.get("missions", {}))
    missions[spec.code] = bloc_v7_data(spec, echantillon, attendus)
    data["missions"] = dict(sorted(missions.items(), key=lambda kv: (kv[0] != "tronc-commun", kv[0])))
    tmp = V7_DATA.with_suffix(".json.tmp")
    ecrire_json(tmp, data)
    tmp.replace(V7_DATA)


def ecrire_verite_privee(spec: MissionSpec, fichiers: dict[str, tuple[list[str], list[dict]]]) -> Path | None:
    """Vérité terrain brute (CSV) dans le dossier PRIVÉ du formateur, s'il existe."""
    if not FORMATEUR.exists():
        return None
    d = VERITE / spec.code
    if d.exists():
        shutil.rmtree(d)
    for nom, (entete, lignes) in fichiers.items():
        ecrire_csv(d / nom, entete, lignes, sep=",")
    return d
